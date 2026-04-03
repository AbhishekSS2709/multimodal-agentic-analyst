"""Excel (.xlsx) loader for the RAG ingestion pipeline."""

from __future__ import annotations

import logging
import uuid
from pathlib import Path
from typing import List

import openpyxl

from src.ingestion.pdf_loader import Document

logger = logging.getLogger(__name__)


def _generate_doc_id() -> str:
    """Generate a unique document identifier."""
    return uuid.uuid4().hex[:12]


def _cell_value(cell) -> str:
    """Return the string representation of a cell value."""
    return str(cell.value) if cell.value is not None else ""


def load_excel(file_path: str) -> List[Document]:
    """Load an Excel workbook and return one Document per sheet.

    Each Document contains a formatted text representation:
        Sheet: <name>
        Columns: ColA | ColB | ...
        ColA: val1 | ColB: val2 | ...   (one line per data row)

    Args:
        file_path: Absolute or relative path to the .xlsx file.

    Returns:
        A list of Document objects, one per sheet.

    Raises:
        FileNotFoundError: If *file_path* does not exist.
    """
    path = Path(file_path)
    if not path.exists():
        raise FileNotFoundError(f"File not found: {file_path}")

    wb = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
    documents: List[Document] = []

    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]

        rows = list(ws.iter_rows())
        if not rows:
            logger.debug("Sheet '%s' is empty, skipping.", sheet_name)
            continue

        # First row is treated as the header
        headers = [_cell_value(cell) for cell in rows[0]]
        data_rows = rows[1:]

        column_count = len(headers)
        row_count = len(data_rows)

        # Build formatted text
        lines: List[str] = [
            f"Sheet: {sheet_name}",
            f"Columns: {' | '.join(headers)}",
        ]
        for row in data_rows:
            cells = [_cell_value(c) for c in row]
            row_str = " | ".join(
                f"{headers[i]}: {cells[i]}" if i < len(cells) else f"{headers[i]}: "
                for i in range(column_count)
            )
            lines.append(row_str)

        full_text = "\n".join(lines)

        document = Document(
            doc_id=_generate_doc_id(),
            text=full_text,
            metadata={
                "source": path.name,
                "file_type": "xlsx",
                "sheet_name": sheet_name,
                "row_count": row_count,
                "column_count": column_count,
                "modality": "table",
            },
            source=str(path),
        )
        documents.append(document)
        logger.debug(
            "Loaded sheet '%s': %d rows x %d cols", sheet_name, row_count, column_count
        )

    wb.close()
    logger.info("Loaded %d sheet(s) from Excel: %s", len(documents), path.name)
    return documents

"""Word document (.docx) loader for the RAG ingestion pipeline."""

from __future__ import annotations

import logging
import uuid
from pathlib import Path
from typing import List

from docx import Document as DocxDocument

from src.ingestion.pdf_loader import Document

logger = logging.getLogger(__name__)


def _generate_doc_id() -> str:
    """Generate a unique document identifier."""
    return uuid.uuid4().hex[:12]


def load_docx(file_path: str) -> List[Document]:
    """Load a Word document and return a single Document with all text.

    Extracts text from all paragraphs and table cells in the document.

    Args:
        file_path: Absolute or relative path to the .docx file.

    Returns:
        A list containing one Document with the full document text.

    Raises:
        FileNotFoundError: If *file_path* does not exist.
    """
    path = Path(file_path)
    if not path.exists():
        raise FileNotFoundError(f"File not found: {file_path}")

    doc = DocxDocument(str(path))

    # Collect paragraph text (includes headings, normal paragraphs, list items)
    paragraph_texts: List[str] = []
    for para in doc.paragraphs:
        stripped = para.text.strip()
        if stripped:
            paragraph_texts.append(stripped)

    # Collect table cell text
    table_texts: List[str] = []
    for table in doc.tables:
        for row in table.rows:
            row_cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if row_cells:
                table_texts.append(" | ".join(row_cells))

    all_parts = paragraph_texts + table_texts
    full_text = "\n".join(all_parts)

    paragraph_count = len(doc.paragraphs)

    document = Document(
        doc_id=_generate_doc_id(),
        text=full_text,
        metadata={
            "source": path.name,
            "file_type": "docx",
            "paragraph_count": paragraph_count,
        },
        source=str(path),
    )

    logger.info("Loaded docx: %s (%d paragraphs)", path.name, paragraph_count)
    return [document]

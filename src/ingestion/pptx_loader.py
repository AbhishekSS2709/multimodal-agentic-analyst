"""PowerPoint (.pptx) loader for the RAG ingestion pipeline."""

from __future__ import annotations

import logging
import uuid
from pathlib import Path
from typing import List

from pptx import Presentation

from src.ingestion.pdf_loader import Document

logger = logging.getLogger(__name__)


def _generate_doc_id() -> str:
    """Generate a unique document identifier."""
    return uuid.uuid4().hex[:12]


def _extract_shape_text(shape) -> str:
    """Return all text from a shape, or empty string if it has none."""
    if shape.has_text_frame:
        return shape.text_frame.text.strip()
    return ""


def _extract_notes_text(slide) -> str:
    """Return speaker-notes text for a slide, or empty string if none."""
    try:
        notes_slide = slide.notes_slide
        tf = notes_slide.notes_text_frame
        return tf.text.strip() if tf else ""
    except Exception:  # noqa: BLE001
        return ""


def load_pptx(file_path: str) -> List[Document]:
    """Load a PowerPoint file and return one Document per slide.

    Each Document contains the slide title, body text, and speaker notes.

    Args:
        file_path: Absolute or relative path to the .pptx file.

    Returns:
        A list of Document objects, one per slide.

    Raises:
        FileNotFoundError: If *file_path* does not exist.
    """
    path = Path(file_path)
    if not path.exists():
        raise FileNotFoundError(f"File not found: {file_path}")

    prs = Presentation(str(path))
    total_slides = len(prs.slides)
    documents: List[Document] = []

    for slide_idx, slide in enumerate(prs.slides, start=1):
        # Collect all shape text on the slide
        shape_texts: List[str] = []
        for shape in slide.shapes:
            text = _extract_shape_text(shape)
            if text:
                shape_texts.append(text)

        # Collect speaker notes
        notes_text = _extract_notes_text(slide)
        has_notes = bool(notes_text)

        parts = shape_texts[:]
        if notes_text:
            parts.append(f"[Notes] {notes_text}")

        full_text = "\n".join(parts)

        document = Document(
            doc_id=_generate_doc_id(),
            text=full_text,
            metadata={
                "source": path.name,
                "file_type": "pptx",
                "slide_number": slide_idx,
                "total_slides": total_slides,
                "has_notes": has_notes,
            },
            source=str(path),
        )
        documents.append(document)
        logger.debug("Loaded slide %d/%d from %s", slide_idx, total_slides, path.name)

    logger.info("Loaded %d slide(s) from pptx: %s", total_slides, path.name)
    return documents

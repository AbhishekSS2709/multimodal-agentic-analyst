"""Tests for src/ingestion/pptx_loader.py — written before implementation (TDD)."""

from __future__ import annotations

import pytest
from pptx import Presentation
from pptx.util import Inches

from src.ingestion.pptx_loader import load_pptx
from src.ingestion.pdf_loader import Document


# ---------------------------------------------------------------------------
# Fixtures
# ---------------------------------------------------------------------------

@pytest.fixture()
def sample_pptx(tmp_path):
    """Create a .pptx with 2 slides, each with a title, body, and speaker notes."""
    pptx_path = tmp_path / "sample.pptx"
    prs = Presentation()

    # Slide 1
    slide_layout = prs.slide_layouts[1]  # Title and Content layout
    slide1 = prs.slides.add_slide(slide_layout)
    slide1.shapes.title.text = "Slide One Title"
    slide1.placeholders[1].text = "Slide one body text."
    notes1 = slide1.notes_slide
    notes1.notes_text_frame.text = "Speaker notes for slide one."

    # Slide 2
    slide2 = prs.slides.add_slide(slide_layout)
    slide2.shapes.title.text = "Slide Two Title"
    slide2.placeholders[1].text = "Slide two body text."
    notes2 = slide2.notes_slide
    notes2.notes_text_frame.text = "Speaker notes for slide two."

    prs.save(str(pptx_path))
    return str(pptx_path)


# ---------------------------------------------------------------------------
# Tests
# ---------------------------------------------------------------------------

class TestLoadPptxReturnsDocuments:
    """load_pptx returns one Document per slide."""

    def test_load_pptx_returns_documents(self, sample_pptx):
        result = load_pptx(sample_pptx)
        assert isinstance(result, list)
        assert len(result) == 2
        assert all(isinstance(doc, Document) for doc in result)


class TestLoadPptxExtractsText:
    """load_pptx captures title, body, and speaker notes text."""

    def test_load_pptx_extracts_text(self, sample_pptx):
        result = load_pptx(sample_pptx)
        texts = [doc.text for doc in result]

        assert any("Slide One Title" in t for t in texts)
        assert any("Slide one body text." in t for t in texts)
        assert any("Speaker notes for slide one." in t for t in texts)

        assert any("Slide Two Title" in t for t in texts)
        assert any("Slide two body text." in t for t in texts)
        assert any("Speaker notes for slide two." in t for t in texts)


class TestLoadPptxMetadata:
    """load_pptx sets required metadata fields for each slide document."""

    def test_load_pptx_metadata(self, sample_pptx):
        result = load_pptx(sample_pptx)
        assert len(result) == 2

        for idx, doc in enumerate(result, start=1):
            meta = doc.metadata
            assert meta["file_type"] == "pptx"
            assert "source" in meta
            assert meta["slide_number"] == idx
            assert meta["total_slides"] == 2
            assert isinstance(meta["has_notes"], bool)
            assert meta["has_notes"] is True

        # source field on the Document dataclass
        assert result[0].source == sample_pptx


class TestLoadNonexistentPptxRaises:
    """load_pptx raises FileNotFoundError for missing files."""

    def test_load_nonexistent_raises(self):
        with pytest.raises(FileNotFoundError):
            load_pptx("/nonexistent/path/missing.pptx")

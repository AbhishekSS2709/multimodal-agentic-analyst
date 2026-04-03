"""Tests for src/ingestion/docx_loader.py — written before implementation (TDD)."""

from __future__ import annotations

import os
import tempfile

import pytest
from docx import Document as DocxDocument

from src.ingestion.docx_loader import load_docx
from src.ingestion.pdf_loader import Document


# ---------------------------------------------------------------------------
# Fixtures
# ---------------------------------------------------------------------------

@pytest.fixture()
def sample_docx(tmp_path):
    """Create a .docx with a heading and two body paragraphs."""
    docx_path = tmp_path / "sample.docx"
    doc = DocxDocument()
    doc.add_heading("Test Heading", level=1)
    doc.add_paragraph("First paragraph content.")
    doc.add_paragraph("Second paragraph content.")
    doc.save(str(docx_path))
    return str(docx_path)


# ---------------------------------------------------------------------------
# Tests
# ---------------------------------------------------------------------------

class TestLoadDocxReturnsDocuments:
    """load_docx returns a non-empty list of Document objects."""

    def test_load_docx_returns_documents(self, sample_docx):
        result = load_docx(sample_docx)
        assert isinstance(result, list)
        assert len(result) >= 1
        assert all(isinstance(doc, Document) for doc in result)


class TestLoadDocxExtractsText:
    """load_docx captures heading and paragraph text."""

    def test_load_docx_extracts_text(self, sample_docx):
        result = load_docx(sample_docx)
        combined = " ".join(doc.text for doc in result)
        assert "Test Heading" in combined
        assert "First paragraph content." in combined
        assert "Second paragraph content." in combined


class TestLoadDocxMetadata:
    """load_docx sets required metadata fields correctly."""

    def test_load_docx_metadata(self, sample_docx):
        result = load_docx(sample_docx)
        assert len(result) == 1
        doc = result[0]

        # source field on the Document dataclass
        assert doc.source == sample_docx

        # metadata dict
        meta = doc.metadata
        assert meta["file_type"] == "docx"
        assert "source" in meta
        assert "paragraph_count" in meta
        assert isinstance(meta["paragraph_count"], int)
        assert meta["paragraph_count"] >= 3  # heading + 2 paragraphs


class TestLoadNonexistentDocxRaises:
    """load_docx raises FileNotFoundError for missing files."""

    def test_load_nonexistent_raises(self):
        with pytest.raises(FileNotFoundError):
            load_docx("/nonexistent/path/missing.docx")
